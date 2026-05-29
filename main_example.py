"""
main_example.py

Reference implementation of the core processing pipeline.
Sensitive values (plant codes, WH mappings, paths) are replaced with generic placeholders.
Split into main.py for production use.
"""

import pandas as pd
from dataclasses import dataclass, field
from pathlib import Path

from utils.date_utils import previous_week
from utils.log_utils import setup_logger
from email_pipeline import email_setup as cfg

OSD_DIR     = cfg.OSD_DIR
IBD_DIR     = cfg.IBD_DIR
MISSING_DIR = Path.home() / 'Desktop' / 'your-project' / 'report'
LOG_DIR     = Path.home() / 'Desktop' / 'your-project' / 'logs'

logger, _log_file = setup_logger(LOG_DIR, 'main')

# Warehouse code → Plant code mapping
WH_TO_PLANT = {
    'WH1': 'PLANT1',
    'WH2': 'PLANT2',
    'WH3': 'PLANT3',
}
PLANT_CODES = list(WH_TO_PLANT.values())

MISSING_FILENAME = "{week} {code} Missing items from OSD Report.xlsx"
IBD_GROUP_COLS   = ['WH N.', 'Week', 'CTN #', 'Serial number', 'Material']

# Fixed 13-column schema for all OSD report files
OSD_COLS = ['Plant', 'Week', 'Date', 'Container', 'SN', 'Live/Drop',
            'Seal in Tact', 'O/S/D', 'Model', 'SN QTY', 'Qty',
            'Serial Number(s)', 'WH Note']


# ============================================================
# Load
# ============================================================

def load_osd_damage_rows(osd_dir: Path, week: str, report: 'ReportResult') -> dict[str, pd.DataFrame]:
    """Read OSD report files and extract Damage (D) rows per plant."""
    result = {}
    for code in PLANT_CODES:
        matches = list(osd_dir.glob(f"*{week}*{code}*"))
        if not matches:
            logger.warning(f"No OSD file found for {code} {week}")
            report.file_missing.append(code)
            continue

        frames = []
        fail_reason = None
        for match in matches:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(match, read_only=True)
                sheet_name = next((s for s in wb.sheetnames if 'summary' in s.strip().lower()), None)
                wb.close()
                if sheet_name is None:
                    fail_reason = f"No 'Summary' sheet found in {match.name}"
                    logger.warning(fail_reason)
                    continue
                raw = pd.read_excel(match, sheet_name=sheet_name, header=None, nrows=10)
                header_row = next((i for i, row in raw.iterrows() if 'O/S/D' in row.values), None)
                if header_row is None:
                    fail_reason = f"Header row not found in {match.name}"
                    logger.warning(fail_reason)
                    continue
                df = pd.read_excel(match, sheet_name=sheet_name, header=header_row)
            except Exception:
                fail_reason = f"Failed to read {match.name}"
                logger.warning(fail_reason)
                continue

            # Normalize columns to fixed OSD_COLS schema regardless of case/variant names
            col_lower = {str(c).strip().lower(): c for c in df.columns}
            rename_map = {}
            for target in OSD_COLS:
                tl = target.lower()
                if tl in col_lower:
                    rename_map[col_lower[tl]] = target
                elif target == 'Serial Number(s)':
                    found = next((c for cl, c in col_lower.items() if 'serial number' in cl), None)
                    if found:
                        rename_map[found] = target
            df = df.rename(columns=rename_map)
            df = df[[c for c in OSD_COLS if c in df.columns]]

            d_rows = df[df['O/S/D'].astype(str).str.strip().str.upper().isin(['D', 'DAMAGE'])].copy()
            frames.append(d_rows)

        if not frames:
            report.fail.append((code, fail_reason or 'Unknown error'))
            continue

        combined = pd.concat(frames, ignore_index=True)
        logger.info(f"{code}: {len(combined)} D rows from {len(frames)} file(s)")
        result[code] = combined
    return result


def load_ibd_origin(ibd_dir: Path, week: str) -> pd.DataFrame:
    """Load the IBD Origin file for the given week."""
    wk = week.replace('W', 'WK')
    matches = list(ibd_dir.glob(f"*IBD Origin*{wk}*"))
    if not matches:
        raise FileNotFoundError(f"IBD Origin file not found for {wk}")
    df = pd.read_excel(matches[0], engine='pyxlsb') if matches[0].suffix == '.xlsb' else pd.read_excel(matches[0])
    logger.info(f"IBD Origin loaded: {len(df)} rows")
    return df


# ============================================================
# Column helpers
# ============================================================

def _sn_col(df):
    return next((c for c in df.columns if 'serial' in str(c).lower()), None)

def _material_col(df):
    return next((c for c in df.columns if 'material' in str(c).lower()), None)

def _model_col(df):
    return next((c for c in df.columns if 'model' in str(c).lower()), None)

def _qty_col_osd(df):
    return next((c for c in df.columns if str(c).strip().lower() in ('qty', 'damage qty')), None)

def _qty_col_ibd(df):
    return next((c for c in df.columns if 'quantity' in str(c).lower()), None)

def _to_int(val, default=1) -> int:
    try:
        return int(val)
    except (ValueError, TypeError):
        return default


# ============================================================
# Core matching
# ============================================================

def _match_and_consume(serial_num: str, material: str, src_qty: int,
                       osd_work: pd.DataFrame,
                       serial_num_col_osd, mod_col_osd, qty_col_osd) -> tuple[int, pd.DataFrame]:
    """
    Match one IBD row against the OSD work set and consume matched rows.
    Serial Number(s) → 1:1 match. No serial number → Model + Qty deduction.
    Returns (remaining_qty, updated_osd_work).
    """
    if serial_num and serial_num_col_osd:
        matched_idx = osd_work[osd_work[serial_num_col_osd].astype(str).str.strip() == serial_num].index.tolist()
        if matched_idx:
            osd_work = osd_work.drop(index=matched_idx).reset_index(drop=True)
            return 0, osd_work

    if material and mod_col_osd and qty_col_osd:
        matched_idx = osd_work[osd_work[mod_col_osd].astype(str).str.strip() == material].index.tolist()
        if not matched_idx:
            return src_qty, osd_work
        osd_qty = sum(_to_int(osd_work.loc[i, qty_col_osd]) for i in matched_idx)
        osd_work = osd_work.drop(index=matched_idx).reset_index(drop=True)
        return max(src_qty - osd_qty, 0), osd_work

    return 0, osd_work


# ============================================================
# IBD vs OSD comparison
# ============================================================

def find_new_missing(ibd: pd.DataFrame, osd_work: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    Compare IBD rows against OSD D-rows.
    Returns (missing_df, remaining_osd_work).
    remaining_osd_work = OSD rows not matched by any IBD row → saved as OVG.
    """
    if ibd.empty:
        return pd.DataFrame(), osd_work

    serial_num_col_ibd = _sn_col(ibd)
    mat_ibd            = _material_col(ibd)
    serial_num_col_osd = _sn_col(osd_work)
    mod_osd            = _model_col(osd_work)
    qty_ibd            = _qty_col_ibd(ibd)
    qty_osd            = _qty_col_osd(osd_work)

    missing_rows = []
    for _, row in ibd.iterrows():
        serial_num = str(row[serial_num_col_ibd]).strip() if serial_num_col_ibd else ''
        serial_num = '' if serial_num.lower() in ('nan', '') else serial_num
        material   = str(row[mat_ibd]).strip() if mat_ibd else ''
        src_qty    = _to_int(row[qty_ibd]) if qty_ibd else 1

        remaining, osd_work = _match_and_consume(
            serial_num, material, src_qty, osd_work,
            serial_num_col_osd, mod_osd, qty_osd
        )
        if remaining > 0:
            new_row = row.copy()
            if qty_ibd:
                new_row[qty_ibd] = remaining
            missing_rows.append(new_row)

    if not missing_rows:
        return pd.DataFrame(), osd_work

    df = pd.DataFrame(missing_rows)
    col_map    = {str(c).strip().lower(): c for c in df.columns}
    qty_col    = col_map.get('quantity')
    group_cols = [col_map[c.lower()] for c in IBD_GROUP_COLS if c.lower() in col_map]
    inbd_col   = col_map.get('inb d/o')

    if qty_col and group_cols:
        agg = {qty_col: 'sum'}
        if inbd_col:
            agg[inbd_col] = 'first'
        df = df.groupby(group_cols, dropna=False).agg(agg).reset_index()
        rename_map = {qty_col: 'Total'}
        if inbd_col:
            rename_map[inbd_col] = 'SN'
        df = df.rename(columns=rename_map)
    else:
        df = df[group_cols]

    ordered = [c for c in ['WH N.', 'Week', 'CTN #', 'SN', 'Material', 'Total', 'Serial number'] if c in df.columns]
    return df[ordered], osd_work


# ============================================================
# OVG
# ============================================================

OVG_FILE       = cfg.OVG_DIR / 'OSD_OVG.xlsx'
KEEP_OVG_WEEKS = 15


def load_ovg_for_plant(plant: str) -> pd.DataFrame:
    if not OVG_FILE.exists():
        return pd.DataFrame(columns=OSD_COLS)
    df = pd.read_excel(OVG_FILE)
    col = next((c for c in df.columns if c.strip().lower() == 'plant'), None)
    if not col:
        return pd.DataFrame(columns=OSD_COLS)
    return df[df[col].astype(str).str.strip() == plant].reset_index(drop=True)


def check_missing_against_ovg(missing: pd.DataFrame, ovg: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Deduct OVG-matched items from the missing list."""
    if missing.empty or ovg.empty:
        return missing, ovg

    serial_num_col_miss = _sn_col(missing)
    mat_miss            = _material_col(missing)
    qty_miss            = next((c for c in missing.columns if str(c).strip().lower() == 'total'), None)
    serial_num_col_ovg  = _sn_col(ovg)
    mod_ovg             = _model_col(ovg)
    qty_ovg             = _qty_col_osd(ovg)

    still_missing = []
    for _, row in missing.iterrows():
        serial_num = str(row[serial_num_col_miss]).strip() if serial_num_col_miss else ''
        serial_num = '' if serial_num.lower() in ('nan', '') else serial_num
        material   = str(row[mat_miss]).strip() if mat_miss else ''
        src_qty    = _to_int(row[qty_miss]) if qty_miss else 1

        remaining, ovg = _match_and_consume(
            serial_num, material, src_qty, ovg,
            serial_num_col_ovg, mod_ovg, qty_ovg
        )
        if remaining > 0:
            new_row = row.copy()
            if qty_miss:
                new_row[qty_miss] = remaining
            still_missing.append(new_row)

    final_missing = pd.DataFrame(still_missing) if still_missing else pd.DataFrame(columns=missing.columns)
    return final_missing, ovg


def save_ovg(remaining_ovg: pd.DataFrame, ovg_new: pd.DataFrame, plant: str, week: str) -> None:
    """Append new OVG rows, remove rows older than KEEP_OVG_WEEKS."""
    new_rows = ovg_new[OSD_COLS].copy() if not ovg_new.empty else pd.DataFrame(columns=OSD_COLS)
    plant_updated = pd.concat([remaining_ovg, new_rows], ignore_index=True)

    cfg.OVG_DIR.mkdir(parents=True, exist_ok=True)

    if OVG_FILE.exists():
        existing = pd.read_excel(OVG_FILE)
        plant_col = next((c for c in existing.columns if c.strip().lower() == 'plant'), None)
        other = existing[existing[plant_col].astype(str).str.strip() != plant] if plant_col else existing
        combined = pd.concat([other, plant_updated], ignore_index=True)
    else:
        combined = plant_updated

    if not combined.empty and 'Week' in combined.columns:
        current_week_num = int(week.replace('W', ''))
        combined = combined[combined['Week'].apply(
            lambda w: (current_week_num - int(w)) % 52 <= KEEP_OVG_WEEKS
        )]

    combined.to_excel(OVG_FILE, index=False)
    logger.info(f"OVG updated: {len(new_rows)} new, {len(remaining_ovg)} remaining for {plant} {week}")


def save_missing(df: pd.DataFrame, code: str, week: str) -> None:
    week_dir = MISSING_DIR / week
    week_dir.mkdir(parents=True, exist_ok=True)
    path = week_dir / MISSING_FILENAME.format(week=week, code=code)
    df.to_excel(path, index=False)
    logger.info(f"Saved: {path.name}")


@dataclass
class ReportResult:
    success:      list[tuple[str, bool]] = field(default_factory=list)
    file_missing: list[str]              = field(default_factory=list)
    fail:         list[tuple[str, str]]  = field(default_factory=list)


# ============================================================
# Run
# ============================================================

def run() -> tuple[ReportResult, str]:
    result = ReportResult()
    week = previous_week()
    logger.info(f"Target week: {week}")

    osd_by_plant = load_osd_damage_rows(OSD_DIR, week, result)
    if not osd_by_plant:
        logger.error("No OSD files found, exiting")
        return result, week

    try:
        ibd = load_ibd_origin(IBD_DIR, week)
    except FileNotFoundError:
        logger.error("IBD Origin file not found")
        result.file_missing.append('IBD')
        return result, week

    wh_col = next((c for c in ibd.columns if str(c).strip().lower() == 'wh n.'), None)

    for code, osd in osd_by_plant.items():
        try:
            wh = next((k for k, v in WH_TO_PLANT.items() if v == code), None)
            ibd_plant = ibd[ibd[wh_col].astype(str).str.strip() == wh] if wh_col and wh else ibd

            osd_work = osd.copy()

            final, ovg_new = find_new_missing(ibd_plant, osd_work)
            logger.info(f"{code}: {len(final)} missing candidates, {len(ovg_new)} OVG new")

            existing_ovg = load_ovg_for_plant(code)
            final, remaining_ovg = check_missing_against_ovg(final, existing_ovg)
            logger.info(f"{code}: {len(final)} final missing after OVG check")

            if not final.empty:
                save_missing(final, code, week)
                result.success.append((code, True))
            else:
                logger.info(f"{code}: no missing items, skipping save")
                result.success.append((code, False))

            save_ovg(remaining_ovg, ovg_new, code, week)

        except Exception:
            logger.exception(f"{code}: unexpected error")
            result.fail.append((code, 'Unexpected error'))

    accounted = {c for c, _ in result.success} | set(result.file_missing) | {c for c, _ in result.fail}
    result.fail += [(c, 'Not processed') for c in PLANT_CODES if c not in accounted]

    logger.info(f"Success: {[c for c, _ in result.success]}")
    logger.info(f"File Missing: {result.file_missing}")
    logger.info(f"Fail: {[(c, r) for c, r in result.fail]}")
    return result, week


if __name__ == '__main__':
    from email_pipeline.sender import load_sender, send_status_report, send_plant_report, send_main_failure

    sender = load_sender()

    week = None
    try:
        result, week = run()

        if sender:
            send_status_report(sender, result, week)

            for code, has_report in result.success:
                if has_report:
                    report_path = MISSING_DIR / week / MISSING_FILENAME.format(week=week, code=code)
                    send_plant_report(sender, code, report_path, week)

    except Exception:
        logger.exception("Unexpected error")
        if sender:
            send_main_failure(sender, _log_file, week or '')
